"""Word (docx) 导出。

与 HTML 模板解耦：按区块顺序生成结构化 Word 文档，字号/行距遵循
SPEC 5.1 的排版参数，保证 Word 版与 PDF 版观感一致。
"""
from __future__ import annotations

import io
from typing import Any

from ..engine import typo
from ..schema import normalize_document


def _pt(v: float) -> float:
    return v


def build_docx(doc: dict[str, Any]) -> io.BytesIO:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = normalize_document(doc)
    design = doc["design"]
    scale = design["fontScale"]

    # 字号（pt）：与 HTML 令牌一致
    FS_BODY = 10.5 * scale
    FS_NAME = 17 * scale
    FS_H2 = 12.5 * scale
    FS_H3 = 11 * scale
    FS_SMALL = 9.5 * scale

    ACCENT = design["accent"].lstrip("#")
    accent_rgb = RGBColor(int(ACCENT[0:2], 16), int(ACCENT[2:4], 16), int(ACCENT[4:6], 16))
    text_rgb = RGBColor(0x1F, 0x29, 0x37)
    soft_rgb = RGBColor(0x37, 0x41, 0x51)
    mute_rgb = RGBColor(0x6B, 0x72, 0x80)

    FONT = "Noto Sans SC" if design["fontFamily"] == "sans" else "Noto Serif SC"
    FONT_LATIN = "Calibri" if design["fontFamily"] == "sans" else "Times New Roman"

    d = Document()

    # 页面设置：A4 + 页边距
    sec = d.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(design["pageMargin"] / 10.0)
    sec.bottom_margin = Cm(design["pageMargin"] / 10.0)
    sec.left_margin = Cm(design["pageMargin"] / 10.0)
    sec.right_margin = Cm(design["pageMargin"] / 10.0)

    # Normal 样式
    normal = d.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(FS_BODY)
    normal.font.color.rgb = text_rgb
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    pf = normal.paragraph_format
    pf.line_spacing = design["lineHeight"]
    pf.space_after = Pt(0)
    pf.space_before = Pt(0)

    def _set_font(run, size=None, bold=None, color=None, latin=None, eastasia=None):
        run.font.name = latin or FONT_LATIN
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        rfonts.set(qn("w:eastAsia"), eastasia or FONT)
        if size is not None:
            run.font.size = Pt(size)
        if bold is not None:
            run.font.bold = bold
        if color is not None:
            run.font.color.rgb = color

    def _para(space_before=0, space_after=0, align=None):
        p = d.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = design["lineHeight"]
        if align is not None:
            p.alignment = align
        return p

    def _bullet(text, size=FS_BODY, color=soft_rgb):
        p = d.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = design["lineHeight"]
        p.paragraph_format.left_indent = Cm(0.5)
        r = p.add_run(typo.tidy(text))
        _set_font(r, size=size, color=color)
        return p

    def _bottom_border(p, color="E5E7EB", sz=6):
        pPr = p._element.get_or_add_pPr()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), str(sz))
        bottom.set(qn("w:space"), "2")
        bottom.set(qn("w:color"), color)
        pbdr.append(bottom)
        pPr.append(pbdr)

    gap_pt = design["sectionGap"] * 0.75  # px -> pt 近似
    content = doc["content"]

    # ---------------- 按 sections 顺序输出 ----------------
    for s in doc["sections"]:
        if not s.get("visible", True):
            continue
        key = s["key"]
        stype = s["type"]

        if key == "profile":
            p = content.get("profile") or {}
            name = typo.tidy(str(p.get("name") or ""))
            if not name:
                continue
            para = _para(align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_font(para.add_run(name), size=FS_NAME, bold=True, color=text_rgb)
            role = typo.tidy(str(p.get("title") or ""))
            if role:
                para = _para(align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_font(para.add_run(role), size=FS_H3, color=accent_rgb)
            contacts = [
                v for k, v in (
                    ("phone", p.get("phone")), ("email", p.get("email")),
                    ("location", p.get("location")), ("age", p.get("age")),
                    ("gender", p.get("gender")), ("desiredSalary", p.get("desiredSalary")),
                    ("availableDate", p.get("availableDate")),
                ) if typo.tidy(str(v or ""))
            ]
            if contacts:
                para = _para(align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_font(para.add_run(typo.join_contact(contacts)), size=FS_SMALL, color=mute_rgb)
            summary = typo.split_lines(str(p.get("summary") or ""))
            if summary:
                _para(space_before=6)
                for ln in summary:
                    para = _para(space_after=2, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                    _set_font(para.add_run(typo.tidy(ln)), color=soft_rgb)
            continue

        if key == "skills":
            sk = content.get("skills") or {}
            featured = sk.get("featuredSkills") or []
            others = sk.get("descriptions") or []
            if not featured and not others:
                continue
            para = _para(space_before=gap_pt, space_after=3)
            _set_font(para.add_run(typo.tidy(s["title"])), size=FS_H2, bold=True, color=text_rgb)
            _bottom_border(para)
            for fs in featured:
                if not isinstance(fs, dict):
                    continue
                nm = typo.tidy(str(fs.get("skill") or ""))
                if nm:
                    p2 = _para(space_after=1)
                    _set_font(p2.add_run(nm), bold=True, color=text_rgb)
                    try:
                        rating = max(0, min(5, int(float(fs.get("rating") or 0))))
                    except (TypeError, ValueError):
                        rating = 0
                    if rating:
                        _set_font(p2.add_run("  " + "★" * rating + "☆" * (5 - rating)),
                                  color=accent_rgb, size=FS_SMALL)
            if others:
                p2 = _para(space_before=2)
                _set_font(p2.add_run(typo.join_contact(others, "、")), color=soft_rgb)
            continue

        if key in ("selfEvaluation", "custom") or stype == "simple":
            val = content.get(key)
            items = val.get("descriptions") if isinstance(val, dict) else val
            items = items if isinstance(items, list) else []
            items = [typo.tidy(str(x)) for x in items if typo.tidy(str(x))]
            if not items:
                continue
            para = _para(space_before=gap_pt, space_after=3)
            _set_font(para.add_run(typo.tidy(s["title"])), size=FS_H2, bold=True, color=text_rgb)
            _bottom_border(para)
            for it in items:
                _bullet(it)
            continue

        # ---------------- 列表型（内置 + 自定义） ----------------
        items = content.get(key)
        if not isinstance(items, list) or not items:
            continue
        fields = [f for f in s.get("fields", []) if isinstance(f, dict) and f.get("key")]
        title_field = next(
            (f["key"] for f in fields
             if f.get("key") in ("name", "company", "school", "project", "title",
                                 "certificate", "item", "organization")),
            None,
        )
        para = _para(space_before=gap_pt, space_after=3)
        _set_font(para.add_run(typo.tidy(s["title"])), size=FS_H2, bold=True, color=text_rgb)
        _bottom_border(para)

        for item in items:
            if not isinstance(item, dict):
                continue
            # 头部行：主标题 | 副标题（左） + 日期（右，用制表位对齐）
            tvals = [typo.tidy(str(item.get(f["key"]) or "")) for f in fields[:2]]
            head_left = " | ".join(v for v in tvals if v)
            date = typo.normalize_date(str(item.get("date") or ""))
            p2 = _para(space_before=4, space_after=1)
            if head_left:
                _set_font(p2.add_run(head_left), size=FS_H3, bold=True, color=text_rgb)
            if date:
                # 右对齐制表位
                tab_stops = p2.paragraph_format.tab_stops
                usable = sec.page_width - sec.left_margin - sec.right_margin
                tab_stops.add_tab_stop(usable, WD_TAB_ALIGNMENT.RIGHT)
                _set_font(p2.add_run("\t" + date), size=FS_SMALL, color=mute_rgb)
            # 其余字段
            for f in fields:
                fk = f["key"]
                if fk in (title_field, "date"):
                    continue
                val = item.get(fk)
                if isinstance(val, list):
                    for v in val:
                        v = typo.tidy(str(v))
                        if v:
                            _bullet(v)
                else:
                    v = typo.tidy("" if val is None else str(val))
                    if v:
                        _bullet(f"{f.get('label') or fk}：{v}")

    buf = io.BytesIO()
    d.save(buf)
    buf.seek(0)
    return buf
